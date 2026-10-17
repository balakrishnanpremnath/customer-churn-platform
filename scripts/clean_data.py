import csv
import json
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

INPUT = ROOT / "data" / "raw" / "online_retail_II.xlsx"
OUTPUT = ROOT / "data" / "processed" / "purchases_clean.csv"
REPORT = ROOT / "data" / "processed" / "cleaning_report.json"


def customer_id_as_text(value):
    if value is None:
        return ""

    text = str(value).strip()

    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]

    return text


def positive_decimal(value):
    try:
        number = Decimal(str(value))

        if number.is_finite() and number > 0:
            return number

    except (InvalidOperation, ValueError, TypeError):
        pass

    return None


def clean_dataset():
    if not INPUT.is_file():
        raise FileNotFoundError(INPUT)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)

    counts = Counter()

    columns = [
        "invoice",
        "stock_code",
        "description",
        "quantity",
        "invoice_date",
        "unit_price",
        "customer_id",
        "country",
        "line_revenue",
    ]

    workbook = load_workbook(
        INPUT,
        read_only=True,
        data_only=True,
    )

    try:
        with OUTPUT.open(
            "w", newline="", encoding="utf-8"
        ) as file:

            writer = csv.DictWriter(
                file,
                fieldnames=columns,
            )

            writer.writeheader()

            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows)

                for row in rows:
                    counts["total_rows"] += 1

                    record = dict(zip(headers, row))

                    invoice = str(
                        record["Invoice"] or ""
                    ).strip()

                    customer = customer_id_as_text(
                        record["Customer ID"]
                    )

                    quantity = positive_decimal(
                        record["Quantity"]
                    )

                    price = positive_decimal(
                        record["Price"]
                    )

                    transaction_date = record["InvoiceDate"]

                    # Exclude cancellation invoices.
                    if invoice.upper().startswith("C"):
                        counts["excluded_cancellations"] += 1
                        continue

                    # Exclude non-positive or invalid quantities.
                    if quantity is None:
                        counts["excluded_quantity"] += 1
                        continue

                    # Exclude non-positive or invalid prices.
                    if price is None:
                        counts["excluded_price"] += 1
                        continue

                    # Customer-level modeling requires an ID.
                    if not customer:
                        counts["excluded_customer_id"] += 1
                        continue

                    if not isinstance(
                        transaction_date, (date, datetime)
                    ):
                        counts["excluded_date"] += 1
                        continue

                    writer.writerow({
                        "invoice": invoice,
                        "stock_code": record["StockCode"],
                        "description": record["Description"],
                        "quantity": str(quantity),
                        "invoice_date": transaction_date.isoformat(),
                        "unit_price": str(price),
                        "customer_id": customer,
                        "country": record["Country"],
                        "line_revenue": str(quantity * price),
                    })

                    counts["clean_rows"] += 1

    finally:
        workbook.close()

    with REPORT.open("w", encoding="utf-8") as file:
        json.dump(
            dict(counts),
            file,
            indent=4,
        )

    print("=" * 50)
    print("CUSTOMER CHURN - DATA CLEANING")
    print("=" * 50)

    for key, value in counts.items():
        print(f"{key}: {value:,}")

    print(f"\nOutput: {OUTPUT.name}")
    print(f"Report: {REPORT.name}")
    print("Data cleaning completed!")


if __name__ == "__main__":
    clean_dataset()