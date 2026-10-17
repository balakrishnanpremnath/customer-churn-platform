from pathlib import Path
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATASET_PATH = (
    PROJECT_ROOT / "data" / "raw"
    / "online_retail_II.xlsx"
)

def inspect_dataset():
    if not DATASET_PATH.is_file():
        raise FileNotFoundError(
            f"Dataset not found: {DATASET_PATH}"
        )

    print("CUSTOMER CHURN PLATFORM")
    print("=" * 50)
    print(f"Dataset: {DATASET_PATH.name}")

    workbook = load_workbook(
        DATASET_PATH,
        read_only=True,
        data_only=True
    )

    try:
        total_rows = 0

        for sheet in workbook.worksheets:
            headers = next(
                sheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True
                )
            )

            record_count = sheet.max_row - 1
            total_rows += record_count

            print(f"\nSheet: {sheet.title}")
            print(f"Records: {record_count:,}")
            print(f"Columns: {sheet.max_column}")
            print(f"Column names: {list(headers)}")

        print("=" * 50)
        print(f"Total records: {total_rows:,}")
        print("Inspection completed!")

    finally:
        workbook.close()

if __name__ == "__main__":
    inspect_dataset()