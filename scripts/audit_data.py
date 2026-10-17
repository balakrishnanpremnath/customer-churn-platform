from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = PROJECT_ROOT / "data" / "raw" / "online_retail_II.xlsx"

REQUIRED_COLUMNS = [
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
]


def audit_dataset():
    if not DATA_PATH.is_file():
        raise FileNotFoundError(f"Dataset not found: {DATA_PATH}")

    print("=" * 55)
    print("CUSTOMER CHURN PLATFORM - DATA QUALITY AUDIT")
    print("=" * 55)

    workbook = load_workbook(
        DATA_PATH,
        read_only=True,
        data_only=True,
    )

    stats = Counter()
    missing = Counter()
    customers = set()
    countries = Counter()

    first_date = None
    last_date = None

    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)

            columns = {
                name: index
                for index, name in enumerate(headers)
            }

            for name in REQUIRED_COLUMNS:
                if name not in columns:
                    raise ValueError(
                        f"Missing column '{name}' in {sheet.title}"
                    )

            sheet_count = 0

            for row in rows:
                if all(value is None for value in row):
                    continue

                sheet_count += 1
                stats["total_rows"] += 1

                for name in REQUIRED_COLUMNS:
                    value = row[columns[name]]

                    if value is None or (
                        isinstance(value, str)
                        and not value.strip()
                    ):
                        missing[name] += 1

                invoice = row[columns["Invoice"]]
                customer = row[columns["Customer ID"]]
                quantity = row[columns["Quantity"]]
                price = row[columns["Price"]]
                invoice_date = row[columns["InvoiceDate"]]
                country = row[columns["Country"]]

                if customer is not None and str(customer).strip():
                    customers.add(str(customer).strip())

                if (
                    isinstance(invoice, str)
                    and invoice.strip().upper().startswith("C")
                ):
                    stats["cancelled_rows"] += 1

                if isinstance(quantity, (int, float)):
                    if quantity <= 0:
                        stats["non_positive_quantity"] += 1
                elif quantity is not None:
                    stats["invalid_quantity"] += 1

                if isinstance(price, (int, float)):
                    if price <= 0:
                        stats["non_positive_price"] += 1
                elif price is not None:
                    stats["invalid_price"] += 1

                if isinstance(invoice_date, (datetime, date)):
                    current_date = (
                        invoice_date.date()
                        if isinstance(invoice_date, datetime)
                        else invoice_date
                    )

                    if first_date is None or current_date < first_date:
                        first_date = current_date

                    if last_date is None or current_date > last_date:
                        last_date = current_date

                elif invoice_date is not None:
                    stats["invalid_date"] += 1

                if country is not None:
                    countries[str(country).strip()] += 1

            print(f"\n{sheet.title}: {sheet_count:,} rows")

    finally:
        workbook.close()

    print("\n" + "=" * 55)
    print("DATA QUALITY SUMMARY")
    print("=" * 55)

    print(f"Total rows: {stats['total_rows']:,}")
    print(f"Unique customers with IDs: {len(customers):,}")
    print(f"Cancelled transaction rows: {stats['cancelled_rows']:,}")
    print(
        "Non-positive quantity rows:",
        f"{stats['non_positive_quantity']:,}"
    )
    print(
        "Non-positive price rows:",
        f"{stats['non_positive_price']:,}"
    )
    print(f"Invalid quantity rows: {stats['invalid_quantity']:,}")
    print(f"Invalid price rows: {stats['invalid_price']:,}")
    print(f"Invalid date rows: {stats['invalid_date']:,}")
    print(f"First transaction date: {first_date}")
    print(f"Last transaction date: {last_date}")

    print("\nMISSING VALUES")
    for name in REQUIRED_COLUMNS:
        count = missing[name]
        percentage = count / stats["total_rows"] * 100
        print(f"{name}: {count:,} ({percentage:.2f}%)")

    print("\nTOP 5 COUNTRIES BY TRANSACTION ROWS")
    for country, count in countries.most_common(5):
        print(f"{country}: {count:,}")

    print("\nAudit completed successfully!")


if __name__ == "__main__":
    audit_dataset()